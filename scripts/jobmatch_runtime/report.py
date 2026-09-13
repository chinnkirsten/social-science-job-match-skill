"""Evidence-gated, local-only Word export. No visual-review claim is made."""
from datetime import datetime, timezone
from io import BytesIO
import os
from pathlib import Path
import tempfile
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from validate_jobs import validate


MODE_FIELDS = {
    'internship': ('earliest_start', 'days_per_week', 'duration_months'),
    'campus_full_time': ('graduation_cohort', 'recruitment_batch', 'graduate_eligibility'),
    'experienced_full_time': ('experience_requirement', 'earliest_start'),
}
LABELS = {
    'internship': '实习', 'campus_full_time': '校招全职',
    'experienced_full_time': '社招全职', 'earliest_start': '最早到岗',
    'days_per_week': '每周天数', 'duration_months': '持续月数',
    'graduation_cohort': '毕业届别', 'recruitment_batch': '招聘批次',
    'graduate_eligibility': '应届资格', 'experience_requirement': '工作年限要求',
    'ready': '可使用（已核对简历或补充说明）', 'verify_first': '需先核实，不可直接使用',
    'create_first': '需先完成，不可写成已有经历',
}
DISPLAY = {
    'open': '正在招聘', 'closed': '已结束招聘', 'unknown': '待确认',
    'pass': '符合已知必需条件', 'fail': '尚未满足必需条件',
    'met': '已有材料支持', 'unmet': '尚未满足',
    'completed': '已完成', 'ongoing': '进行中', 'planned': '计划中',
    'unconfirmed': '待确认', 'employer_official': '雇主官网',
    'official_ats': '官方招聘系统', 'employer_verified_platform': '已核实雇主身份的直发平台',
    'official_repost': '机构转载', 'aggregator': '招聘信息汇总网站',
    'sample_observation': '仅说明本次参考岗位的共同要求', 'none': '暂无可比较的岗位',
    'SourceScout': '岗位信息整理', 'EvidenceMapper': '经历匹配与简历建议', 'Auditor': '独立复核',
}


def _display(value):
    """Translate controlled labels only; never rewrite quoted source material."""
    if type(value) is bool:
        return '是' if value else '否'
    return DISPLAY.get(str(value), str(value))


def _filled(value):
    return isinstance(value, str) and bool(value.strip())


def _strings(value):
    return isinstance(value, list) and bool(value) and all(_filled(x) for x in value)


def _extensions(data, stage=False):
    """Validate presentation content separately from the factual gate."""
    errors = []
    if not _filled(data.get('report_summary')):
        errors.append('report_summary must be a non-empty string')
    selected = {j['id'] for j in data['jobs'] if j.get('selected') is True
                or (stage and j.get('proposed_selection') is True)}
    variants = data.get('resume_variants')
    ids, covered = {}, set()
    if not isinstance(variants, list) or not variants:
        errors.append('resume_variants must be a non-empty list')
        variants = []
    if len(variants) > 3:
        errors.append('Consolidate resume_variants into at most three evidence-based structures')
    for index, item in enumerate(variants):
        prefix = f'resume_variants[{index}]'
        if not isinstance(item, dict):
            errors.append(prefix + ' must be an object')
            continue
        if not all(_filled(item.get(k)) for k in ('id', 'name')):
            errors.append(prefix + ' requires id and name')
        if not _strings(item.get('changes')):
            errors.append(prefix + ' requires concrete changes')
        if not _strings(item.get('job_ids')) or not set(item['job_ids']) <= selected:
            errors.append(prefix + ' requires selected job_ids')
        else:
            covered.update(item['job_ids'])
        if _filled(item.get('id')):
            if item['id'] in ids:
                errors.append(prefix + ' duplicate variant id')
            ids[item['id']] = item
    if covered != selected:
        errors.append('resume_variants must cover every selected job')
    actions = data.get('action_plan')
    if not isinstance(actions, list) or not actions:
        errors.append('action_plan must be a non-empty list')
        actions = []
    covered = set()
    for index, item in enumerate(actions):
        prefix = f'action_plan[{index}]'
        if not isinstance(item, dict):
            errors.append(prefix + ' must be an object')
            continue
        job_id, variant_id = item.get('job_id'), item.get('resume_variant_id')
        if not isinstance(job_id, str) or job_id not in selected:
            errors.append(prefix + ' job_id must identify a selected job')
        else:
            covered.add(job_id)
        if not _filled(item.get('action')) or not _strings(item.get('materials')):
            errors.append(prefix + ' requires action and materials')
        if not isinstance(variant_id, str) or variant_id not in ids:
            errors.append(prefix + ' requires an existing resume_variant_id')
        elif not _strings(ids[variant_id].get('job_ids')) or job_id not in ids[variant_id]['job_ids']:
            errors.append(prefix + ' variant does not cover this job')
    if covered != selected:
        errors.append('action_plan must cover every selected job')
    for job in data['jobs']:
        if job['id'] not in selected:
            continue
        details = job.get('details')
        if not isinstance(details, dict):
            errors.append(f"{job['id']}: details missing")
            continue
        for field in ('salary', 'deadline'):
            if not _filled(details.get(field)):
                errors.append(f"{job['id']}: details.{field} missing; explicitly state undisclosed if needed")
        fields = details.get('mode_fields')
        for field in MODE_FIELDS[data['employment_mode']]:
            if not isinstance(fields, dict) or not _filled(fields.get(field)):
                errors.append(f"{job['id']}: details.mode_fields.{field} missing")
    return errors


def _inspect(path):
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with ZipFile(path) as archive:
        if archive.testzip():
            raise ValueError('Word ZIP integrity failure')
        root = ET.fromstring(archive.read('word/document.xml'))
        if root.findall('.//w:shd', ns):
            raise ValueError('Word contains forbidden shading')
        for name in archive.namelist():
            if name.startswith('word/') and name.endswith('.xml'):
                tree = ET.fromstring(archive.read(name))
                for color in tree.findall('.//w:color', ns):
                    if color.get('{%s}val' % ns['w']) not in ('000000', 'auto'):
                        raise ValueError('Word contains non-black text')
        if not root.findall('.//w:bookmarkStart', ns):
            raise ValueError('Word navigation bookmarks missing')
    return {'ooxml_valid': True, 'visual_review': 'pending'}


def _monochrome(path):
    """Also clean the template's unused stylesWithEffects/numbering defaults."""
    from lxml import etree
    namespace = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    with ZipFile(path) as archive:
        entries = [(entry, archive.read(entry.filename)) for entry in archive.infolist()]
    buffer = BytesIO()
    with ZipFile(buffer, 'w') as archive:
        for entry, content in entries:
            if entry.filename.startswith('word/') and entry.filename.endswith('.xml'):
                tree = etree.fromstring(content)
                for shade in tree.iter('{%s}shd' % namespace):
                    shade.getparent().remove(shade)
                for border in list(tree.iter('{%s}pBdr' % namespace)):
                    border.getparent().remove(border)
                for color in tree.iter('{%s}color' % namespace):
                    color.attrib.clear()
                    color.set('{%s}val' % namespace, '000000')
                content = etree.tostring(tree, encoding='UTF-8', xml_declaration=True, standalone=True)
            archive.writestr(entry, content)
    Path(path).write_bytes(buffer.getvalue())


def render_report(data: dict, output: Path, now=None, stage=False) -> dict:
    """Write a new DOCX, never overwrite. Validation failures return diagnostics.

    Additional formal fields: report_summary (string), resume_variants
    [{id,name,job_ids,changes}], action_plan
    [{job_id,action,materials,resume_variant_id}], and selected jobs' details
    {salary,deadline,mode_fields}. MODE_FIELDS defines mode-specific keys.
    A stage export may omit presentation fields but cannot bypass evidence errors.
    A successful export remains render_pending until external visual review.
    """
    output = Path(output)
    if output.suffix.lower() != '.docx':
        return {'ok': False, 'status': 'blocked', 'errors': ['output must have .docx extension']}
    if output.exists() or output.is_symlink():
        return {'ok': False, 'status': 'blocked', 'errors': ['output already exists; overwrite forbidden']}
    try:
        checked = validate(data, now)
    except (TypeError, ValueError, KeyError) as exc:
        return {'ok': False, 'status': 'blocked', 'errors': ['invalid input structure: ' + str(exc)]}
    if not checked.get('evidence_ok', checked['ok']) or (not stage and not checked['ok']):
        return {'ok': False, 'status': 'blocked', 'errors': checked['errors'], 'validation': checked}
    missing = _extensions(data, stage)
    if missing and not stage:
        return {'ok': False, 'status': 'blocked', 'errors': missing, 'validation': checked}
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
    except ImportError:
        return {'ok': False, 'status': 'blocked', 'errors': ['Install report extra: python-docx is required']}

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    document.core_properties.author = 'Job Match Skill'
    document.core_properties.title = '求职岗位与简历调整报告'
    document.core_properties.subject = '岗位匹配与简历建议；资料检查通过，排版待检查'
    for style in document.styles:
        if style.type == 1 or style.type == 2:
            style.font.name = 'Arial'
            style.font.size = Pt(12)
            style.font.color.rgb = RGBColor(0, 0, 0)
            if style.element.rPr is not None:
                fonts = style.element.rPr.rFonts
                if fonts is not None:
                    fonts.set(qn('w:eastAsia'), '宋体')
        for node in list(style.element.iter(qn('w:shd'))):
            node.getparent().remove(node)
        for node in list(style.element.iter(qn('w:color'))):
            node.attrib.clear()
            node.set(qn('w:val'), '000000')
    document.styles['Normal'].paragraph_format.space_after = Pt(6)
    document.styles['Normal'].paragraph_format.line_spacing = 1.3
    for name, size in [('Title', 22), ('Heading 1', 17), ('Heading 2', 14), ('Heading 3', 12)]:
        document.styles[name].font.size = Pt(size)
        document.styles[name].font.bold = True
    footer = section.footer.paragraphs[0]
    footer.add_run('第 ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    footer.add_run(' 页 · 排版待检查')

    def paragraph(text, style=None):
        return document.add_paragraph(str(text), style)

    def link(p, text, destination=None, anchor=None):
        element = OxmlElement('w:hyperlink')
        if destination:
            element.set(qn('r:id'), p.part.relate_to(destination, RT.HYPERLINK, is_external=True))
        if anchor:
            element.set(qn('w:anchor'), anchor)
        run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '000000'); props.append(color)
        size = OxmlElement('w:sz'); size.set(qn('w:val'), '24'); props.append(size)
        underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
        run.append(props); node = OxmlElement('w:t'); node.text = str(text); run.append(node)
        element.append(run); p._p.append(element)

    bookmark_counter = 0

    def heading(text, anchor, level=1):
        nonlocal bookmark_counter
        bookmark_counter += 1
        p = document.add_heading(text, level)
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(bookmark_counter)); start.set(qn('w:name'), anchor)
        end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(bookmark_counter))
        p._p.insert(0, start); p._p.append(end)

    def table(headers, rows):
        tab = document.add_table(rows=1, cols=len(headers))
        tab.autofit = False
        widths = [17 / len(headers)] * len(headers)
        for column, width in zip(tab.columns, widths):
            column.width = Cm(width)
        for cell, value in zip(tab.rows[0].cells, headers):
            cell.text = str(value)
        repeat = OxmlElement('w:tblHeader'); tab.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for cell, value in zip(tab.add_row().cells, row):
                cell.text = str(value)
        for row in tab.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(10.5)
                        run.font.color.rgb = RGBColor(0, 0, 0)
        borders = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            edge = OxmlElement('w:' + side)
            for key, value in [('val', 'single'), ('sz', '4'), ('color', '000000')]:
                edge.set(qn('w:' + key), value)
            borders.append(edge)
        tab._tbl.tblPr.append(borders)
        return tab

    ledger = {item['id']: item for item in data['corpus']['candidate_evidence']}
    corpus = {item['corpus_id']: item for item in data['corpus']['records']}

    def refs(items):
        return '\n'.join(f"经历 {ref}：{ledger[ref]['source_id']} / {ledger[ref]['locator']}；"
                         f"{_display(ledger[ref]['state'])}；已确认：{_display(ledger[ref]['confirmed'])}" for ref in items)

    selected = [job for job in data['jobs'] if job.get('selected')]
    proposed = [job for job in data['jobs'] if stage and data.get('runtime_version') is not None and not job.get('selected')
                and job.get('proposed_selection') is True]
    detailed = sorted(selected + proposed, key=lambda job: (job.get('rank_position', 10**9), job['id']))
    paragraph('求职岗位与简历调整报告', 'Title')
    paragraph('阶段报告' if stage else '求职报告')
    paragraph('本报告由程序整理。资料检查不能代替逐项核实，排版仍需检查。')
    if data.get('synthetic') is True:
        paragraph('模拟测试数据，仅用于软件验证；不代表真实候选人或可投岗位。')
    directions = data['corpus'].get('target_directions', [])
    paragraph(f"招聘模式：{LABELS[data['employment_mode']]}；方向：{' / '.join(directions) or '未单独设置'}；地区：{data['corpus']['geography']}")
    paragraph(f"生成时间：{(now or datetime.now(timezone.utc)).isoformat()}")
    paragraph(f"主清单：{checked['companies']} 家 / 目标 {checked['target']} 家；{len(selected)} 个岗位。")
    if proposed:
        paragraph(f'机器分析草稿／未确认可投：另有 {len(proposed)} 个待人工复核岗位，不计入已确认主清单。')
    sections = [('报告摘要', 'summary'), ('样本与来源', 'sources'), ('岗位与逐条调整', 'jobs'),
                ('简历版本与投递安排', 'actions'), ('待确认与未推荐岗位', 'appendix'), ('经历记录与材料出处', 'ledger')]
    paragraph('目录', 'Heading 1')
    for label, anchor in sections:
        link(paragraph(''), label, anchor=anchor)
    for index, job in enumerate(detailed):
        link(paragraph(''), f"{index + 1}. {job['company']} · {job['title']}", anchor=f'job_{index}')
    heading('报告摘要', 'summary')
    paragraph(data.get('report_summary') or '摘要尚未提供；本阶段稿不作为完整交付。')
    action_by_job = {a.get('job_id'): a for a in data.get('action_plan', []) if isinstance(a, dict)}
    if detailed:
        paragraph('优先准备的岗位（机器草稿仍须先复核）' if stage else '优先准备的岗位', 'Heading 2')
        for job in detailed[:3]:
            action = action_by_job.get(job['id'], {})
            paragraph(f"{job['company']} · {job['title']}；简历版本：{action.get('resume_variant_id', '待整理')}")
            paragraph('下一步：' + str(action.get('action', '先核对来源与资格')))
            paragraph('材料：' + '；'.join(action.get('materials', []) if _strings(action.get('materials')) else ['待确认']))
            paragraph('准备量：' + str(action.get('effort_estimate', '未估计，需根据材料缺口确认')))
    if stage:
        paragraph('阶段限制：' + '\n'.join(checked['errors'] + missing or ['本次明确按阶段稿导出。']))
    heading('样本与来源', 'sources')
    basis = data['evidence_basis']
    for field, title in [('method_version', '分析方法版本'), ('corpus_as_of', '岗位信息采集截至'),
                         ('candidate_evidence_count', '经历记录数量'), ('market_corpus_count', '可比较的完整岗位数量'),
                         ('market_employer_count', '涉及雇主数量'), ('trend_claim_level', '分析适用范围')]:
        paragraph(f'{title}：{_display(basis[field])}')
    paragraph('这些资料只反映本次参考岗位的情况，不能代表整个招聘市场，也不能用于预测录用概率。')
    coverage = data['corpus'].get('search_coverage', {})
    if coverage:
        paragraph(f"岗位发现：{coverage.get('method')}；发现线索 {coverage.get('candidate_links', 0)} 条；"
                  f"人工选中具体来源 {coverage.get('reviewed_source_links', 0)} 条；"
                  f"结果截断：{_display(coverage.get('truncated', False))}。")
    execution = data.get('execution', {})
    calls = execution.get('model_calls', [])
    if calls:
        paragraph('本次自动分析步骤：' + '、'.join(_display(c.get('role', 'unknown')) for c in calls))
    for adapter in execution.get('adapters', []):
        paragraph(f"本次使用工具：{adapter.get('component_id')}；版本：{adapter.get('upstream_version')}；使用时间：{adapter.get('invoked_at')}；复用已保存结果：{_display(adapter.get('cache_hit'))}")
    for warning in execution.get('warnings', []):
        paragraph(f"补充工具未运行成功：{warning.get('component_id')}；岗位编号：{warning.get('source_id')}。{warning.get('impact', '')}")
    for source in basis['market_sources']:
        p = paragraph(f"{source['name']} | {_display(source['source_tier'])} | {source['records']} 条 | ")
        link(p, '来源', source['url'])
    for taxonomy in basis.get('taxonomies', []):
        p = paragraph(f"{taxonomy['name']} {taxonomy['version']}：{taxonomy['use']} | ")
        link(p, '分类体系', taxonomy['url'])
    heading('岗位与逐条调整', 'jobs')
    for index, job in enumerate(detailed):
        heading(f"{index + 1}. {job['company']} · {job['title']}", f'job_{index}', 2)
        if not job.get('selected'):
            paragraph('机器分析草稿／未确认可投。以下是待人工复核的分析，不是投递推荐。')
        row = corpus[job['corpus_id']]
        paragraph(f"编号：{job['id']}；投递顺序：第 {job.get('rank_position', '待排')} 位；优先级：{job['priority']}；方向：{job.get('target_direction', '未单独归类')}；地点：{' / '.join(row['locations'])}")
        basis = job.get('rank_basis', {})
        if basis:
            paragraph(f"排序依据：有材料支持的对应 {basis.get('direct_resume_mappings', 0)} 项；"
                      f"未解决加分项 {basis.get('unresolved_preferred_conditions', 0)} 项；"
                      f"来源：{_display(basis.get('source_tier'))}。该顺序不是面试或录用概率。")
        paragraph(f"招聘状态：{_display(job['status'])}；申请条件：{_display(job['eligibility'])}；核查时间：{job['checked_at']}")
        paragraph(f"信息来源：{_display(job['source_tier'])}；在招依据：{job['open_evidence']}")
        link(paragraph(''), '具体 JD', job['jd_url'])
        if job.get('apply_url'):
            link(paragraph(''), '投递入口', job['apply_url'])
        if job.get('application_method'):
            paragraph('投递方式：' + job['application_method'])
        details = job.get('details') if isinstance(job.get('details'), dict) else {}
        paragraph('薪酬：' + str(details.get('salary', '缺少信息，待补齐')))
        paragraph('截止时间：' + str(details.get('deadline', '缺少信息，待补齐')))
        mode_fields = details.get('mode_fields') if isinstance(details.get('mode_fields'), dict) else {}
        for field in MODE_FIELDS[data['employment_mode']]:
            paragraph(f"{LABELS[field]}：{mode_fields.get(field, '缺少信息，待补齐')}")
        paragraph('JD 职责', 'Heading 3')
        for value in row['responsibilities']:
            paragraph(value)
        paragraph('岗位要求与申请条件', 'Heading 3')
        for requirement in job['requirements']:
            paragraph(f"{requirement['text']} | {'必需' if requirement['required'] else '优先'} | {_display(requirement['result'])}")
            paragraph('JD 原文：' + requirement['jd_evidence'])
            paragraph('相关经历：' + requirement['candidate_evidence'])
            paragraph(refs(requirement.get('evidence_refs', [])) or '尚未找到已确认的相关经历。')
        paragraph('逐项要求对照', 'Heading 3')
        for mapping in job['mappings']:
            table(['项目', '具体说明'], [('岗位要求', mapping['requirement']), ('相关经历', mapping['resume_evidence']),
                  ('材料出处', refs(mapping.get('evidence_refs', [])) or '尚无对应经历记录'),
                  ('待补充之处', mapping['gap']), ('修改建议', mapping['action'])])
        paragraph('简历改写示例与放置位置', 'Heading 3')
        for rewrite in job['rewrites']:
            paragraph(f"位置：{rewrite['placement']}；使用状态：{LABELS[rewrite['use_status']]}")
            paragraph(rewrite['text'])
            paragraph(refs(rewrite['evidence_refs']))
        paragraph(f"岗位资料编号：{row['corpus_id']}；采集时间：{row['captured_at']}；内容校验码（SHA256）：{row['source_sha256']}")
        paragraph('招聘原文', 'Heading 3')
        paragraph(row['source_text'])
    heading('简历版本与投递安排', 'actions')
    for variant in data.get('resume_variants', []) if isinstance(data.get('resume_variants'), list) else []:
        if not isinstance(variant, dict):
            continue
        paragraph(f"{variant.get('id', '字段缺失')} · {variant.get('name', '字段缺失')}", 'Heading 2')
        paragraph('适用岗位：' + ', '.join(variant.get('job_ids', []) if _strings(variant.get('job_ids')) else []))
        for change in variant.get('changes', []) if _strings(variant.get('changes')) else []:
            paragraph(change)
    for action in data.get('action_plan', []) if isinstance(data.get('action_plan'), list) else []:
        if not isinstance(action, dict):
            continue
        paragraph(f"岗位 {action.get('job_id', '字段缺失')}；简历版本 {action.get('resume_variant_id', '字段缺失')}")
        paragraph(action.get('action', '字段缺失'))
        paragraph('材料：' + '；'.join(action.get('materials', []) if _strings(action.get('materials')) else []))
    heading('待确认与未推荐岗位', 'appendix')
    pending = [job for job in data['jobs'] if not job.get('selected') and job not in proposed]
    if not pending:
        paragraph('本次输入没有待确认或排除记录；这不表示全网没有其他岗位。')
    for job in pending:
        paragraph(f"{job['id']} · {job['company']} · {job['title']}", 'Heading 2')
        paragraph(f"招聘状态：{_display(job['status'])}；申请条件：{_display(job['eligibility'])}；原因：{job['reason']}")
        link(paragraph(''), '待确认或排除岗位 JD', job['jd_url'])
        for req in job['requirements']:
            paragraph(f"{req['text']} | {_display(req['result'])} | {req['jd_evidence']} | {req['candidate_evidence']}")
            paragraph(refs(req.get('evidence_refs', [])) or '没有已确认引用。')
    represented = {job.get('corpus_id') for job in data['jobs'] if isinstance(job.get('corpus_id'), str)}
    for row in data['corpus']['records']:
        if row['corpus_id'] in represented:
            continue
        paragraph(f"未列入推荐的岗位资料：{row['corpus_id']} · {row['title']}", 'Heading 2')
        paragraph(f"招聘状态：{_display(row['status'])}；纳入同批比较：{_display(row['comparable'])}；已取得完整 JD：{_display(row['full_jd'])}")
        paragraph('记录原因：' + (row.get('excluded_reason') or '本次未生成对应岗位分析，不作为可投结论。'))
        link(paragraph(''), '招聘页面', row['jd_url'])
    heading('经历记录与材料出处', 'ledger')
    for item in ledger.values():
        paragraph(f"{item['id']} | {item['source_id']} | {item['locator']}", 'Heading 2')
        paragraph(f"经历状态：{_display(item['state'])}；已确认：{_display(item['confirmed'])}")
        paragraph(item['text'])
    try:
        output.parent.mkdir(parents=True, exist_ok=True)
        fd, temporary = tempfile.mkstemp(prefix='.jobmatch-', suffix='.docx', dir=output.parent)
        os.close(fd)
    except OSError as exc:
        return {'ok': False, 'status': 'blocked', 'errors': [str(exc)], 'validation': checked}
    try:
        document.save(temporary)
        _monochrome(temporary)
        inspection = _inspect(temporary)
        os.link(temporary, output)  # Atomic create, fails even if destination appears concurrently.
    except (OSError, ValueError) as exc:
        return {'ok': False, 'status': 'blocked', 'errors': [str(exc)], 'validation': checked}
    finally:
        Path(temporary).unlink(missing_ok=True)
    return {'ok': True, 'status': 'render_pending', 'output': str(output), 'stage': stage,
            'companies': checked['companies'], 'selected_jobs': len(selected), 'proposed_jobs': len(proposed),
            'diagnostics': checked['errors'] + missing, 'validation': checked, **inspection}
