"""Real local parsing and synthetic discovery contracts; no live-job claims."""
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import prepare_job_match
from jobmatch_runtime import preparation as prep
from jobmatch_runtime.cache import Cache
from jobmatch_runtime.common import AdapterError


class PreparationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.resume = self.root / 'synthetic.txt'
        self.resume.write_text('Synthetic fixture, not a real candidate\nAccounting student, graduating 2027\nPlanned Excel course\n', encoding='utf-8')
        self.intake = {'employment_mode': 'campus_full_time', 'locations': ['杭州'], 'recruitment_season': 'autumn'}

    def review_resume(self, draft):
        return {'draft_sha256': draft['draft_sha256'], 'human_confirmed': True, 'intake_confirmed': True,
                'records': [{'id': r['id'], 'keep': i > 0, 'state': 'ongoing' if i == 1 else 'planned',
                             'text_verified': True} for i, r in enumerate(draft['proposed_records'])]}

    def source_draft(self):
        page = {'url': 'https://careers.example.test/', 'captured_at': '2026-09-13T00:00:00Z', 'cache_hit': False,
                'links': ['https://careers.example.test/jobs/1', 'https://careers.example.test/jobs/1?utm_source=test',
                          'https://other.example.test/jobs/2', 'https://careers.example.test/about',
                          'https://careers.example.test/jobs/3']}
        config = {**self.intake, 'allowed_source_hosts': ['careers.example.test'],
                  'discovery_url_contains': ['/jobs/'], 'discovery_max_links': 1}
        with patch.object(prep.sources, 'fetch', return_value=page):
            return prep.discover_links([{'url': page['url']}], config, Cache(self.root / 'cache'))

    def source_review(self, draft):
        return {'draft_sha256': draft['draft_sha256'], 'human_confirmed': True,
                'sources': [{'id': draft['candidate_links'][0]['id'], 'specific_jd_confirmed': True,
                             'scope_confirmed': True, 'source_tier': 'employer_official'}]}

    def test_mode_required_before_read_or_model(self):
        with patch.object(prep, 'parse_local') as parser:
            with self.assertRaises(AdapterError) as exc:
                prep.prepare_resume(self.resume, {'locations': ['杭州']})
            self.assertEqual(exc.exception.code, 'mode_required')
            parser.assert_not_called()

    def test_inline_credentials_rejected_before_export(self):
        with self.assertRaises(AdapterError) as exc:
            prep.prepared_bundle({'preparation': {'reviewed': True}}, self.intake, [],
                                 {'llm': {'api_key': 'synthetic-not-a-key'}})
        self.assertEqual(exc.exception.code, 'inline_credential')

    def test_all_three_modes_without_degree_inference(self):
        for mode in prep.MODES:
            result = prep.prepare_resume(self.resume, {'employment_mode': mode, 'locations': ['杭州']})
            self.assertEqual(result['intake']['employment_mode'], mode)
            self.assertNotIn('candidate_evidence', result)
            self.assertTrue(all(r['confirmed'] is False and r['state'] == 'unconfirmed' for r in result['proposed_records']))

    def test_autumn_internship_conflict_rejected(self):
        with self.assertRaises(AdapterError) as exc:
            prep.validate_intake({**self.intake, 'employment_mode': 'internship'})
        self.assertEqual(exc.exception.code, 'season_mode_conflict')

    def test_native_docx_reads_paragraph_and_table(self):
        from docx import Document
        path = self.root / 'synthetic.docx'
        document = Document()
        document.add_paragraph('Synthetic accounting education')
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'Planned course'
        document.save(path)
        result = prep.parse_local(path)
        self.assertEqual([r['text'] for r in result['records']], ['Synthetic accounting education', 'Planned course'])

    def test_native_pdf_real_extract(self):
        from reportlab.pdfgen.canvas import Canvas
        path = self.root / 'synthetic.pdf'
        canvas = Canvas(str(path))
        canvas.drawString(40, 750, 'Synthetic accounting student graduating in 2027')
        canvas.save()
        result = prep.parse_local(path)
        self.assertIn('graduating in 2027', result['records'][0]['text'])
        self.assertEqual(result['records'][0]['locator'], 'page 1, line 1')

    def test_blank_pdf_requires_ocr(self):
        from reportlab.pdfgen.canvas import Canvas
        path = self.root / 'blank.pdf'
        canvas = Canvas(str(path)); canvas.showPage(); canvas.save()
        with self.assertRaises(AdapterError) as exc:
            prep.parse_local(path)
        self.assertEqual(exc.exception.code, 'ocr_required')

    def test_malformed_docx_errors_do_not_expose_paths(self):
        path = self.root / 'synthetic.docx'
        path.write_bytes(b'not an office document')
        with self.assertRaises(AdapterError) as exc:
            prep.parse_local(path)
        self.assertEqual(exc.exception.code, 'resume_parse_failed')
        self.assertNotIn(str(path), str(exc.exception))

    def test_remote_model_needs_consent_no_call(self):
        with patch('jobmatch_runtime.llm.http_json') as http:
            with self.assertRaises(AdapterError) as exc:
                prep.prepare_resume(self.resume, self.intake,
                                    {'llm': {'base_url': 'https://models.example.test/v1', 'model': 'test'}}, use_model=True)
            self.assertEqual(exc.exception.code, 'privacy_consent_required')
            http.assert_not_called()

    def test_model_selection_cannot_drop_or_rewrite(self):
        with patch.object(prep.ModelClient, 'call', return_value={'record_ids': ['exp-0002']}):
            result = prep.prepare_resume(self.resume, self.intake, use_model=True)
        self.assertEqual(len(result['proposed_records']), 3)
        self.assertEqual(result['proposed_records'][1]['text'], 'Accounting student, graduating 2027')
        self.assertTrue(result['proposed_records'][1]['model_suggested_experience'])

    def test_forged_model_ids_rejected(self):
        with patch.object(prep.ModelClient, 'call', return_value={'record_ids': ['fabricated']}):
            with self.assertRaises(AdapterError):
                prep.prepare_resume(self.resume, self.intake, use_model=True)

    def test_confirmation_preserves_planned_state_and_excludes_header(self):
        draft = prep.prepare_resume(self.resume, self.intake)
        candidate, intake = prep.confirm_resume(draft, self.review_resume(draft))
        self.assertEqual(intake['recruitment_season'], 'autumn')
        self.assertEqual([r['state'] for r in candidate['candidate_evidence']], ['ongoing', 'planned'])
        self.assertEqual(len(candidate['candidate_evidence']), 2)
        self.assertNotIn(str(self.root), json.dumps(draft))

    def test_no_confirmation_no_export(self):
        draft = prep.prepare_resume(self.resume, self.intake)
        review = self.review_resume(draft); review['human_confirmed'] = False
        with self.assertRaises(AdapterError) as exc:
            prep.confirm_resume(draft, review)
        self.assertEqual(exc.exception.code, 'human_review_required')

    def test_draft_tampering_invalidates_review(self):
        draft = prep.prepare_resume(self.resume, self.intake)
        review = self.review_resume(draft)
        draft['proposed_records'][1]['text'] = 'Invented experience'
        with self.assertRaises(AdapterError) as exc:
            prep.confirm_resume(draft, review)
        self.assertEqual(exc.exception.code, 'draft_changed')

    def test_missing_record_decision_rejected(self):
        draft = prep.prepare_resume(self.resume, self.intake)
        review = self.review_resume(draft); review['records'].pop()
        with self.assertRaises(AdapterError):
            prep.confirm_resume(draft, review)

    def test_discovery_bounds_scope_dedup_and_not_verified(self):
        draft = self.source_draft()
        self.assertEqual(len(draft['candidate_links']), 1)
        self.assertTrue(draft['truncated'])
        self.assertEqual(draft['candidate_links'][0]['status'], 'unverified_link')
        self.assertNotIn('sources', draft)

    def test_discovery_failure_is_visible_not_dead_link(self):
        with patch.object(prep.sources, 'fetch', side_effect=AdapterError('unsafe_url', 'reserved DNS')):
            draft = prep.discover_links([{'url': 'https://careers.example.test/'}], self.intake, Cache(self.root / 'cache'))
        self.assertEqual(draft['candidate_links'], [])
        self.assertEqual(draft['failures'][0]['code'], 'unsafe_url')
        self.assertEqual(draft['status'], 'needs_human_review')

    def test_source_review_exports_only_selected_specific_jds(self):
        draft = self.source_draft()
        result = prep.confirm_sources(draft, self.source_review(draft))
        self.assertEqual(result, [{'url': 'https://careers.example.test/jobs/1', 'source_tier': 'employer_official'}])
        self.assertNotIn('status', result[0])

    def test_source_scope_review_required(self):
        draft = self.source_draft(); review = self.source_review(draft)
        review['sources'][0]['scope_confirmed'] = False
        with self.assertRaises(AdapterError):
            prep.confirm_sources(draft, review)

    def test_cli_export_compatible_and_private(self):
        draft = prep.prepare_resume(self.resume, self.intake); sources = self.source_draft()
        values = {'resume-draft': draft, 'resume-review': self.review_resume(draft),
                  'source-draft': sources, 'source-review': self.source_review(sources), 'config': {}}
        args = ['export']
        for name, value in values.items():
            path = self.root / (name + '.json'); path.write_text(json.dumps(value), encoding='utf-8')
            args += ['--' + name, str(path)]
        output = self.root / 'prepared'
        args += ['--output-dir', str(output)]
        self.assertEqual(prepare_job_match.main(args), 0)
        self.assertEqual((output.stat().st_mode & 0o777), 0o700)
        for name in ('candidate', 'config', 'sources'):
            self.assertEqual(((output / (name + '.json')).stat().st_mode & 0o777), 0o600)
        self.assertEqual(json.loads((output / 'config.json').read_text())['employment_mode'], 'campus_full_time')
        self.assertEqual(prepare_job_match.main(args), 2)


if __name__ == '__main__':
    unittest.main()
